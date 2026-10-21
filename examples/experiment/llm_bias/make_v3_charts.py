"""Charts + one Excel for the V3 single-addition tests (log 14.45). Laptop only (matplotlib, openpyxl).

    oasis-env/bin/python examples/experiment/llm_bias/make_v3_charts.py [out_dir]     (default ~/Desktop/LLM_Bias/V3)

Every test is shown next to V1 (the `v1` run, same users, seeds).
- posts_per_user_<test>.png: % of users who wrote 0 / 1 / 2 / 3+ posts, per AI.
- actions_<test>.png: every action used, how many times, V1 vs test, per AI (posting) or per reader (reading).
- actions_per_screen_<test>.png: % of screens with 0 / 1 / 2 / 3 / 4+ actions.
- LLM_Bias_V3.xlsx: the same numbers as tables, plus users who used each action at least once.
"""
import json, os, sys
from collections import Counter, defaultdict

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
V3 = os.path.join(HERE, "..", "..", "..", "data", "llm_bias", "v3_tests")
OUT = os.path.expanduser(sys.argv[1] if len(sys.argv) > 1 else "~/Desktop/LLM_Bias/V3")
M = ["gemma3-12b", "llama3.1-8b", "qwen3-8b"]
NAME = {"gemma3-12b": "gemma", "llama3.1-8b": "llama", "qwen3-8b": "qwen"}
COLOR = {"gemma3-12b": "#2a78d6", "llama3.1-8b": "#eb6834", "qwen3-8b": "#1baf7a"}  # dataviz slots 1-3
V1GRAY, INK, INK2, GRID = "#b4b2ab", "#0b0b0b", "#52514e", "#e6e5e1"
POST = {"p_think": "qwen thinks first", "p_menu": "create_post is the only action",
        "p_2a": '+ "This Reddit is brand new ... everyone will read next."',
        "p_2b": "+ \"You're helping build this Reddit's first feed ...\"",
        "p_2c": '+ "It\'s posting time ... read the new posts afterwards."', "p_screen": "the whole V3 posting screen"}
READ = {"r_think": "qwen thinks first", "r_refresh": "refresh works (post comes back)",
        "r_json": "JSON-only off in the engine", "r_menu": "menu flipped, least-used first"}
load = lambda f: [json.loads(l) for l in open(f) if l.strip()]
plt.rcParams.update({"font.family": "Helvetica", "font.size": 10, "text.color": INK, "axes.labelcolor": INK2,
                     "xtick.color": INK2, "ytick.color": INK2, "axes.edgecolor": GRID, "figure.facecolor": "white"})


def posting(t, m):
    return load(f"{V3}/{t}/r101/posting_{m}.jsonl")


def reading(t, reader):  # first-view screens + all rows, every poster's set, noise floor left out
    rs = [r for p in M for r in load(f"{V3}/{t}/r101/reading_{p}__{reader}_a30.jsonl")]
    return rs, sum(r.get("view", 1) == 1 for r in rs)


def bucket(n, top):
    return f"{top}+" if n >= top else str(n)


def style(ax):
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x" if ax.get_xlim() else "y", color=GRID, lw=0.8)
    ax.set_axisbelow(True)


# ---------- numbers
def post_counts(rows):
    c = Counter(bucket(sum(a["action"] == "create_post" for a in r["actions"]), 3) for r in rows)
    return [c[k] for k in ("0", "1", "2", "3+")]


def action_table(rows):  # action -> (times, users who used it)
    times, users = Counter(), defaultdict(set)
    for r in rows:
        for a in r["actions"]:
            times[a["action"]] += 1
            users[a["action"]].add(r["user_id"])
    return {k: (times[k], len(users[k])) for k in times}


def per_screen(rows):
    return Counter(bucket(sum(x["action"] != "do_nothing" for x in r["actions"]), 4)  # do_nothing = 0
                   for r in rows if r.get("view", 1) == 1)


# ---------- charts
def chart_posts(t):
    fig, axs = plt.subplots(1, 3, figsize=(13, 3.8), sharey=True)
    ks = ["0", "1", "2", "3+"]
    for ax, m in zip(axs, M):
        for j, run in enumerate(["v1", t]):
            rows = posting(run, m)
            pc = [100 * v / len(rows) for v in post_counts(rows)]
            x = [k + (j - 0.5) * 0.4 for k in range(4)]
            ax.bar(x, pc, 0.38, color=V1GRAY if run == "v1" else COLOR[m], label="V1" if run == "v1" else t)
            for xx, v in zip(x, pc):
                ax.text(xx, v + 1.5, f"{v:.0f}", ha="center", fontsize=7, color=INK2)
        ax.set_xticks(range(4), ["0 posts", "1 post", "2 posts", "3+ posts"])
        ax.set_title(f"{NAME[m]} wrote the posts", fontsize=10, loc="left")
        ax.legend(frameon=False, fontsize=8)
        ax.set_ylim(0, 108)
        style(ax)
        ax.grid(axis="x", visible=False)
    axs[0].set_ylabel("% of the 100 users")
    fig.suptitle(f"Posts per user: V1 vs {t} ({POST[t]})", x=0.01, ha="left", fontsize=11)
    return save(fig, f"posts_per_user_{t}.png")


def chart_actions(t, rows_of, who, total_of, unit):
    tabs = {m: (action_table(rows_of("v1", m)), action_table(rows_of(t, m))) for m in M}
    names = sorted({a for b, x in tabs.values() for a in list(b) + list(x)},
                   key=lambda a: -sum(b.get(a, (0, 0))[0] + x.get(a, (0, 0))[0] for b, x in tabs.values()))
    fig, axs = plt.subplots(1, 3, figsize=(13, 0.32 * len(names) + 1.6), sharey=True)
    for ax, m in zip(axs, M):
        b, x = tabs[m]
        y = range(len(names))
        bv = [b.get(a, (0, 0))[0] for a in names]
        tv = [x.get(a, (0, 0))[0] for a in names]
        ax.barh([i - 0.2 for i in y], bv, 0.38, color=V1GRAY, label="V1")
        ax.barh([i + 0.2 for i in y], tv, 0.38, color=COLOR[m], label=t)
        top = max(bv + tv + [1])
        for i, (p, q) in enumerate(zip(bv, tv)):
            ax.text(p + top * 0.01, i - 0.2, f"{p:,}", va="center", fontsize=6.5, color=INK2)
            ax.text(q + top * 0.01, i + 0.2, f"{q:,}", va="center", fontsize=6.5, color=INK)
        ax.set_xlim(0, top * 1.22)
        ax.set_title(f"{who} {NAME[m]}  ({total_of('v1', m):,} -> {total_of(t, m):,} {unit})", fontsize=10, loc="left")
        ax.legend(frameon=False, fontsize=8, loc="lower right")
        style(ax)
        ax.grid(axis="y", visible=False)
    axs[0].set_yticks(range(len(names)), names)
    axs[0].invert_yaxis()
    fig.suptitle(f"Actions used (times): V1 vs {t} ({(POST | READ)[t]})", x=0.01, ha="left", fontsize=11)
    return save(fig, f"actions_{t}.png")


def chart_per_screen(t, rows_of, who, unit):
    fig, axs = plt.subplots(1, 3, figsize=(13, 3.6), sharey=True)
    ks = ["0", "1", "2", "3", "4+"]
    for ax, m in zip(axs, M):
        for j, run in enumerate(["v1", t]):
            c = per_screen(rows_of(run, m))
            n = sum(c.values())
            pc = [100 * c[k] / n for k in ks]
            x = [k + (j - 0.5) * 0.4 for k in range(5)]
            ax.bar(x, pc, 0.38, color=V1GRAY if run == "v1" else COLOR[m], label="V1" if run == "v1" else t)
            for xx, v in zip(x, pc):
                if v >= 1:
                    ax.text(xx, v + 1, f"{v:.0f}", ha="center", fontsize=7, color=INK2)
        ax.set_xticks(range(5), ks)
        ax.set_xlabel("actions in one answer (do_nothing counts as 0)")
        ax.set_title(f"{who} {NAME[m]}", fontsize=10, loc="left")
        ax.legend(frameon=False, fontsize=8)
        ax.set_ylim(0, 108)
        style(ax)
        ax.grid(axis="x", visible=False)
    axs[0].set_ylabel(f"% of {unit}")
    fig.suptitle(f"How many actions per answer: V1 vs {t} ({(POST | READ)[t]})", x=0.01, ha="left", fontsize=11)
    return save(fig, f"actions_per_screen_{t}.png")


def save(fig, name):
    fig.tight_layout()
    p = os.path.join(OUT, name)
    fig.savefig(p, dpi=150)
    plt.close(fig)
    return p


# ---------- Excel
def sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(10, min(46, max(len(str(c.value or "")) for c in col) + 2))
    return ws


def main():
    os.makedirs(OUT, exist_ok=True)
    pngs = []
    post_rows = lambda t, m: posting(t, m)
    read_rows = lambda t, m: reading(t, m)[0]
    for t in POST:
        pngs += [chart_posts(t), chart_actions(t, post_rows, "posts by", lambda t_, m: len(posting(t_, m)), "users"),
                 chart_per_screen(t, post_rows, "posting:", "the 100 users")]
    for t in READ:
        pngs += [chart_actions(t, read_rows, "reader", lambda t_, m: reading(t_, m)[1], "screens"),
                 chart_per_screen(t, read_rows, "reader", "posts read")]

    wb = Workbook()
    wb.remove(wb.active)
    rows = []
    for t in ["v1"] + list(POST):
        for m in M:
            rs = posting(t, m)
            c = post_counts(rs)
            rows.append([t, (POST | {"v1": "nothing (V1 again)"})[t], NAME[m], len(rs), *c,
                         *[round(100 * v / len(rs), 1) for v in c], sum(r["outcome"] != "chose" for r in rs),
                         sum(a["action"] == "create_post" for r in rs for a in r["actions"])])
    sheet(wb, "Posts per user", ["test", "the one change", "AI", "users", "0 posts", "1 post", "2 posts", "3+ posts",
                                 "% 0", "% 1", "% 2", "% 3+", "unreadable answers", "posts written"], rows)
    rows = []
    for t in ["v1"] + list(POST):
        for m in M:
            rs = posting(t, m)
            for a, (n, u) in sorted(action_table(rs).items(), key=lambda kv: -kv[1][0]):
                rows.append([t, NAME[m], a, n, u, round(100 * u / len(rs), 1)])
    sheet(wb, "Posting actions", ["test", "AI", "action", "times used", "users who used it", "% of users"], rows)
    rows = []
    for t in ["v1"] + list(READ):
        for m in M:
            for p in M:
                rs = load(f"{V3}/{t}/r101/reading_{p}__{m}_a30.jsonl")
                n = sum(r.get("view", 1) == 1 for r in rs)
                for a, (k, u) in sorted(action_table(rs).items(), key=lambda kv: -kv[1][0]):
                    rows.append([t, NAME[m], NAME[p], a, k, u, n, round(100 * k / n, 1)])
    sheet(wb, "Reading actions", ["test", "reader", "posts by", "action", "times used", "users who used it (of 30)",
                                  "screens", "times per 100 screens"], rows)
    rows = []
    for turn, tests, get in [("posting", ["v1"] + list(POST), post_rows), ("reading", ["v1"] + list(READ), read_rows)]:
        for t in tests:
            for m in M:
                c = per_screen(get(t, m))
                n = sum(c.values())
                ks = ["0", "1", "2", "3", "4+"]
                rows.append([turn, t, NAME[m], n, *[c[k] for k in ks], *[round(100 * c[k] / n, 1) for k in ks]])
    sheet(wb, "Actions per answer", ["turn", "test", "AI", "answers", "0 (incl. do_nothing)", "1", "2", "3", "4+",
                                     "% 0", "% 1", "% 2", "% 3", "% 4+"], rows)
    for p in pngs:
        ws = wb.create_sheet(os.path.basename(p)[:-4][:31])
        img = Image(p)
        img.width, img.height = img.width * 0.55, img.height * 0.55
        ws.add_image(img, "A1")
    x = os.path.join(OUT, "LLM_Bias_V3.xlsx")
    wb.save(x)
    print(f"{len(pngs)} charts + {x}")


if __name__ == "__main__":
    main()
