"""One Excel workbook per round of Test 6: every person's whole scroll, post by post, with both AIs' reactions.

    python make_round_workbooks.py            # writes data/llm_bias/two_ai/by_round/round_01.xlsx … round_15.xlsx

Sheets: README (what every column means), Scrolls (one row per person per post, in the order they saw them),
People (the 100 people, the exact text the AI read, live counts of their reactions), Posts (every post in full,
with live vote counts). Counts are Excel formulas (COUNTIFS on the Scrolls sheet), so they stay right if you filter
or edit; Excel, Numbers and Google Sheets calculate them when the file opens.
"""

import ast
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import authors  # noqa: E402
from topics import TOPICS  # noqa: E402

D = os.path.join(authors.DATA, "two_ai")
OUT = os.path.join(D, "by_round")
A, B = "gemma4:e2b", "gemma3:1b"
ACT = {"like": "upvote", "dislike": "downvote", "nothing": "nothing"}
CARE = {-2: "dislikes it", -1: "bored by it", 0: "doesn't care", 1: "enjoys it", 2: "loves it"}
HABIT = {"generous": "upvotes almost everything", "typical": "typical", "harsh": "hard to please"}
TOPIC_ORDER = ["personal_finance", "cars", "farming", "cooking", "tech"]
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)
HEAD = PatternFill("solid", start_color="DDE3EA")


def short(ai):
    return ai.split(":")[0]


def write_sheet(ws, headers, rows, widths, wrap_cols=()):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = FB, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap_cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    os.makedirs(OUT, exist_ok=True)
    R = pd.read_csv(os.path.join(D, "reactions.csv"))
    P = pd.read_csv(os.path.join(D, "posts.csv"))
    U = pd.read_csv(os.path.join(D, "users.csv"))
    U["aff"] = U.topic_affinity.apply(ast.literal_eval)
    users = U.set_index("id")
    for rnd in sorted(R["round"].unique()):
        V = R[R["round"] == rnd]
        posts = P[(P["round"] == rnd) & P.ok].copy()
        posts["pid"] = posts.apply(lambda p: f"{TOPICS[p.topic]['sub']} #{p.slot_in_topic + 1} by {short(p.author)}", axis=1)
        pid = dict(zip(posts.key, posts.pid))
        shown = set(V.post_key)
        posts = posts[posts.key.isin(shown)]
        g4 = V[V.played_by == A].set_index(["user_id", "post_key"])
        g3 = V[V.played_by == B].set_index(["user_id", "post_key"])
        # the scroll order is the same for both AIs; check it rather than assume it
        both = g4[["topic_rank", "pos_in_topic"]].join(g3[["topic_rank", "pos_in_topic"]], rsuffix="_3")
        assert (both.topic_rank == both.topic_rank_3).all() and (both.pos_in_topic == both.pos_in_topic_3).all()
        pmap = posts.set_index("key")
        rows = []
        for u in sorted(V.user_id.unique()):
            mine = g4.loc[u].sort_values(["topic_rank", "pos_in_topic"])
            p = users.loc[u]
            for n, (key, r) in enumerate(mine.iterrows(), 1):
                post = pmap.loc[key]
                r3 = g3.loc[(u, key)]
                rows.append([int(u), p.realname, int(p.age), p.profession, HABIT[p.voting], n, TOPICS[post.topic]["sub"],
                             int(r.topic_rank) + 1, CARE[int(r.affinity)], pid[key], short(post.author), post.title, post.body,
                             ACT.get(r.action, "(no answer)"), r.reason if isinstance(r.reason, str) else "",
                             ACT.get(r3.action, "(no answer)"), r3.reason if isinstance(r3.reason, str) else "", None])
        wb = Workbook()
        rd = wb.active
        rd.title = "README"
        ws = wb.create_sheet("Scrolls")
        hdr = ["Person #", "Name", "Age", "Job", "Voting habit", "Scroll position (1 = first post they saw)", "Subreddit",
               "Subreddit's place in their scroll (1 = favourite)", "How much they care about this subreddit", "Post ID",
               "Written by", "Post title", "Post text (exactly what they read)", f"When {short(A)} played them",
               f"{short(A)}'s reason", f"When {short(B)} played them", f"{short(B)}'s reason", "Did both AIs do the same?"]
        write_sheet(ws, hdr, rows, [9, 18, 6, 22, 22, 11, 18, 12, 16, 30, 10, 45, 70, 14, 40, 14, 55, 12])
        for i in range(2, len(rows) + 2):
            ws[f"R{i}"] = f'=IF(N{i}=P{i},"same","different")'
            ws[f"R{i}"].font = F
        # People
        wp = wb.create_sheet("People")
        ph = ["Person #", "Name", "Username", "Age", "Gender", "Job", "Lives in", "Voting habit"] + \
             [TOPICS[t]["sub"] for t in TOPIC_ORDER] + \
             [f"{short(A)} played them: upvotes", "downvotes", "nothing", f"{short(B)} played them: upvotes", "downvotes", "nothing",
              f"{short(A)} played them: upvotes on {short(A)}'s posts", f"…on {short(B)}'s posts",
              f"{short(B)} played them: upvotes on {short(B)}'s posts", f"…on {short(A)}'s posts",
              "Posts where both AIs did the same", "Exact description the AI was given"]
        prow = []
        for u in sorted(V.user_id.unique()):
            p = users.loc[u]
            prow.append([int(u), p.realname, p.username, int(p.age), p.gender, p.profession, p.place, HABIT[p.voting]] +
                        [CARE[p.aff[t]] for t in TOPIC_ORDER] + [None] * 11 + [p.persona])
        write_sheet(wp, ph, prow, [9, 18, 16, 6, 10, 22, 22, 20] + [14] * 5 + [12] * 11 + [90])
        S = "Scrolls!"
        for i in range(2, len(prow) + 2):
            a = f"{S}$A:$A,$A{i}"
            f = {"N": f'=COUNTIFS({a},{S}$N:$N,"upvote")', "O": f'=COUNTIFS({a},{S}$N:$N,"downvote")',
                 "P": f'=COUNTIFS({a},{S}$N:$N,"nothing")', "Q": f'=COUNTIFS({a},{S}$P:$P,"upvote")',
                 "R": f'=COUNTIFS({a},{S}$P:$P,"downvote")', "S": f'=COUNTIFS({a},{S}$P:$P,"nothing")',
                 "T": f'=COUNTIFS({a},{S}$K:$K,"{short(A)}",{S}$N:$N,"upvote")',
                 "U": f'=COUNTIFS({a},{S}$K:$K,"{short(B)}",{S}$N:$N,"upvote")',
                 "V": f'=COUNTIFS({a},{S}$K:$K,"{short(B)}",{S}$P:$P,"upvote")',
                 "W": f'=COUNTIFS({a},{S}$K:$K,"{short(A)}",{S}$P:$P,"upvote")',
                 "X": f'=COUNTIFS({a},{S}$R:$R,"same")'}
            for col, fm in f.items():
                wp[f"{col}{i}"] = fm
                wp[f"{col}{i}"].font = F
        # Posts
        wq = wb.create_sheet("Posts")
        qh = ["Post ID", "Subreddit", "Pair # (both AIs wrote one from the same job card)", "Written by", "Job card: subject",
              "Job card: kind of post", "Job card: posting as", "Title", "Words", "Full text",
              f"{short(A)}'s people: upvotes", "downvotes", "nothing", f"{short(B)}'s people: upvotes", "downvotes", "nothing",
              f"Score from {short(A)}'s people (up − down)", f"Score from {short(B)}'s people (up − down)"]
        posts["order"] = posts.topic.map(TOPIC_ORDER.index)
        posts = posts.sort_values(["order", "slot_in_topic", "author"], key=lambda s: s if s.name != "author" else s.map({A: 0, B: 1}))
        qrow = [[p.pid, TOPICS[p.topic]["sub"], int(p.slot_in_topic) + 1, short(p.author)] +
                list(ast.literal_eval(p.brief).values()) + [p.title, int(p.words), p.body] + [None] * 8 for p in posts.itertuples()]
        write_sheet(wq, qh, qrow, [32, 18, 12, 10, 34, 30, 30, 45, 7, 80] + [12] * 8)
        for i in range(2, len(qrow) + 2):
            a = f"{S}$J:$J,$A{i}"
            for col, fm in {"K": f'=COUNTIFS({a},{S}$N:$N,"upvote")', "L": f'=COUNTIFS({a},{S}$N:$N,"downvote")',
                            "M": f'=COUNTIFS({a},{S}$N:$N,"nothing")', "N": f'=COUNTIFS({a},{S}$P:$P,"upvote")',
                            "O": f'=COUNTIFS({a},{S}$P:$P,"downvote")', "P": f'=COUNTIFS({a},{S}$P:$P,"nothing")',
                            "Q": f"=K{i}-L{i}", "R": f"=N{i}-O{i}"}.items():
                wq[f"{col}{i}"] = fm
                wq[f"{col}{i}"].font = F
        # README
        lines = [
            (f"Test 6, round {rnd}: every person's scroll, post by post", FB),
            (f"{V.user_id.nunique()} people × {len(posts)} posts = {len(rows):,} rows on the Scrolls sheet. Each row is one person seeing "
             f"one post; both AIs' reactions are side by side.", F),
            ("", F),
            ("How the round worked", FB),
            (f"Both AIs ({short(A)} and {short(B)}) wrote 5 posts for each of the 5 subreddits from the same job cards. Then {short(A)} "
             f"pretended to be each person and scrolled every post, and later {short(B)} pretended to be the same person and scrolled "
             "the same posts in the same order. For each post the AI answered upvote, downvote or nothing, with a short reason.", F),
            ("The AI playing a person sees only: the person's description (People sheet, last column), the subreddit, the title and "
             "the post text. Never who wrote it, other people's votes or anything it voted before. Every vote is a fresh question.", F),
            ("", F),
            ("Sheets", FB),
            ("Scrolls: one row per person per post, in the exact order that person saw them. Filter 'Person #' to follow one "
             "person through all their posts; filter 'Did both AIs do the same?' to see where the two AIs disagree about the "
             "same person.", F),
            ("People: the 100 people, how much each cares about each subreddit, the exact description the AI was given, and counts "
             "of their reactions this round (formulas over the Scrolls sheet).", F),
            ("Posts: every post of the round in full, its job card, and how each AI's 100 people reacted (formulas).", F),
            ("", F),
            ("Words used", FB),
            ("Scroll position: 1 = the first post the person saw. People see their favourite subreddit first and their least "
             "favourite last; inside a subreddit the posts are shuffled (a fixed shuffle per person).", F),
            ("How much they care: dislikes it / bored by it / doesn't care / enjoys it / loves it (fixed for each person).", F),
            ("Voting habit: upvotes almost everything / typical / hard to please (part of the description the AI reads).", F),
            ("Post ID: subreddit, pair number and author, e.g. 'r/cars #2 by gemma4'. The post with the same pair number by the "
             "other AI was written from the same job card.", F),
            ("'(no answer)': the AI's answer could not be read even after retrying (2 votes in the whole test, both round 13).", F),
            ("", F),
            ("Source: data/llm_bias/two_ai/reactions.csv, posts.csv and users.csv (GitHub KGordo11/oasis, branch llm-bias). "
             "Made by examples/experiment/llm_bias/make_round_workbooks.py.", F),
        ]
        for i, (t, font) in enumerate(lines, 1):
            rd[f"A{i}"] = t
            rd[f"A{i}"].font = Font(name="Arial", size=13, bold=True) if i == 1 else font
            rd[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")
        rd.column_dimensions["A"].width = 130
        wb.calculation.fullCalcOnLoad = True
        path = os.path.join(OUT, f"round_{rnd:02d}.xlsx")
        wb.save(path)
        print(f"round {rnd:2d}: {len(rows):,} rows, {len(posts)} posts, {os.path.getsize(path) // 1024} KB")


if __name__ == "__main__":
    main()
