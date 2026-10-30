"""One Excel workbook for the V3 big sim, built from the cleaned data (v3_clean.py; log 14.60).

    oasis-env/bin/python examples/experiment/llm_bias/make_v3_workbook.py <out.xlsx> data/llm_bias/v3 [data/llm_bias/v3_flipped]

Give part 2's folder too (after `v3_clean.py data/llm_bias/v3_flipped`). Sheets: Read me; 1 No timeline (V1 rounds 101
and 102 next to each part's tick 0, every action); 2 With timeline (each part's whole scroll, every action); 3 Own-AI
bias (every set, 95% ranges); one sheet per part per world, tick by tick, every action; Worlds; Actions; Cleaning;
effort; Posts; Users; Reads and Follow-ups. Numbers are plain values from the clean CSVs (and V1's raw rows).
"""
import csv, os, sys
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import analyze_v2 as A  # noqa: E402
from v3_big_summary import load  # noqa: E402

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT, DIRS = sys.argv[1], sys.argv[2:]
PART = ["part 1", "part 2"][:len(DIRS)]
F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
HEAD, BASE = PatternFill("solid", start_color="DDE3EA"), PatternFill("solid", start_color="FFF4D6")
NUM = lambda v: int(v) if str(v).lstrip("-").isdigit() else float(v) if v.replace(".", "", 1).lstrip("-").isdigit() else v  # noqa: E731


def rows(d, name):
    return [{k: NUM(v) for k, v in r.items()} for r in csv.DictReader(open(os.path.join(d, "clean", name)))]


data = [{n: rows(d, n + ".csv") for n in ("reads", "followups", "posts", "by_cell_tick", "by_cell_action", "action_effort",
                                          "effort_tiers")} for d in DIRS]
WORLDS = sorted({(r["reader"], r["posts_by"]) for r in data[0]["reads"]}, key=lambda c: (c[0] != c[1], c))
wb = Workbook()


def sheet(title, head, body, widths=None, note=None, base_col=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = FB
        r0 = 3
    for j, h in enumerate(head, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.alignment = FB, HEAD, Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(body, r0 + 1):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v).font = F
        if base_col is not None and row[base_col[0]] == row[base_col[1]]:
            for j in range(1, len(row) + 1):
                ws.cell(i, j).fill = BASE
    for j in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 12)
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(head))}{r0 + len(body)}"
    return ws


# Read me
ws = wb.active
ws.title = "Read me"
TEXT = [
    ("V3 big sim, cleaned data", ""),
    ("Parts in this file", ", ".join(f"{p} = {d}" for p, d in zip(PART, DIRS))),
    ("", ""),
    ("WORDS", ""),
    ("world", "One reader AI playing all 100 users, reading the posts one AI wrote. 3 reader AIs x 3 post writers = 9 worlds. "
              "Each world is separate: nothing in one world is seen in another."),
    ("baseline", "A world where the AI reads posts its own AI wrote (gemma reads gemma's posts, ...). Shaded yellow."),
    ("tick", "Everyone scrolls at the same pace. At tick 0 each user reads the first post of their own scroll, at tick 1 the "
             "second, and so on. Users see comments other users left in EARLIER ticks. Tick 0 = nobody has commented yet, like V1."),
    ("post read", "One user reading one post. If the user pressed refresh and the post came back, both looks count as ONE post read."),
    ("actions chosen", "Every action the AI gave in its answer on the post screen, except 'do nothing'. This is total engagement, "
                       "the main measure, counted the same way as V1. Actions on follow-up screens are counted apart."),
    ("actions done", "Only the chosen actions that worked in the world. An action fails when, e.g., it likes a comment ID that does "
                     "not exist, likes the same post twice, or asks to open a second screen while one is open."),
    ("upvoted / downvoted / commented", "Number of posts read where the AI upvoted (like_post) / downvoted / wrote a comment."),
    ("did nothing", "Posts read where the AI's only answer was 'do nothing'."),
    ("comments on screen", "Comments by other users that were on the screen when the post was read (added up over the posts read)."),
    ("follow-up screen", "A screen the user opened from the post: a profile, a search, an opened post, messages. Not a post read; "
                         "kept apart (Follow-ups sheet)."),
    ("per 100 / %", "Divided by posts read. The worlds have different numbers of posts (gemma 99, llama 119, qwen 102 posts, each "
                    "read by 99 users), so compare rates between worlds, counts within one world."),
    ("", ""),
    ("SHEETS", ""),
    ("1 No timeline", "Each world with NO timeline: V1 rounds 101 and 102 (V1 never had one) and each V3 part's tick 0 (each user's "
                      "first post, nobody has commented yet; 100 posts read per world). Every action, count and per 100."),
    ("2 With timeline", "Each world over the WHOLE scroll (every tick), part 1 and part 2 kept apart, then the change. Every action."),
    ("3 Own-AI bias", "Does each AI give its own AI's posts more than the other AI does? Every set, total engagement and upvotes, "
                      "with the 95% range. CLEAR = the range leaves out 0."),
    ("P1 / P2 world sheets", "One sheet per part per world: every tick (round of the scroll) on its own row, then 'all'. "
                             "Counts, rates, and how many times each action was chosen at that tick."),
    ("Cleaning", "What was kept out of the main tables and why, with counts; every failed action and why it failed."),
    ("Worlds", "Each world all together (the 'all' rows of the world sheets side by side)."),
    ("Actions", "Every action, per world: how many times chosen and done, and per 100 posts read."),
    ("Effort tiers / Action effort", "How much each action asks of the AI (one click, copy an ID, write text, write a post, "
                                     "go to another screen first) and how often each AI picks it, how often it fails, how "
                                     "long the answer is and how long it takes. Every chosen action on every screen is in "
                                     "clean/actions.csv (too big for this file)."),
    ("Posts", "Each post: who wrote it, topic, length, and in each world that read it: times read, upvoted, downvoted, commented."),
    ("Users", "Each user in each world: posts read, actions, upvotes, comments, follow-up screens."),
    ("Reads / Follow-ups", "Every post read and every follow-up screen, one row each. Filter these to check any number in the file."),
    ("", ""),
    ("Source", "v3_clean.py writes <part>/clean/*.csv from the raw files in <part>/r101/; make_v3_workbook.py builds this file. "
               "v3_clean.py checks its numbers against the published ones before writing (log 14.53, 14.60)."),
]
for i, (a, b) in enumerate(TEXT, 1):
    ws.cell(i, 1, a).font = Font(name="Arial", size=13, bold=True) if i == 1 else FB
    ws.cell(i, 2, b).font = F
    ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 120

# Cleaning (the markdown report, table by table)
for p, d in zip(PART, DIRS):
    lines = [l.rstrip("\n") for l in open(os.path.join(d, "clean", "cleaning.md"))]
    ws = wb.create_sheet(f"Cleaning {p}" if len(DIRS) > 1 else "Cleaning")
    r = 1
    for l in lines:
        if l.startswith("|---"):
            continue
        cells = [c.strip() for c in l.strip("|").split("|")] if l.startswith("|") else [l.lstrip("# ")]
        for j, c in enumerate(cells, 1):
            ws.cell(r, j, NUM(c.replace(",", "")) if l.startswith("|") else c).font = FB if not l.startswith("|") or r == 1 else F
        r += 1
    for j, w in enumerate((70, 22, 40, 12, 12, 12), 1):
        ws.column_dimensions[get_column_letter(j)].width = w

# the measures for one world, from a by_cell_tick row (V3) or from V1's raw rows
COUNTS = ["posts read", "actions chosen", "actions done", "upvoted", "downvoted", "commented", "did nothing",
          "comments on screen", "follow-up screens"]
RATES = ["engagement per 100", "done per 100", "upvoted %", "commented %"]
tick = [{(r["reader"], r["posts_by"], r["tick"]): r for r in x["by_cell_tick"]} for x in data]
ACTS = sorted({k for tk in tick for r in tk.values() for k in r} - set(COUNTS + RATES) - {"reader", "posts_by", "tick"})
S = {"gemma3:12b": "gemma", "llama3.1:8b": "llama", "qwen3:8b": "qwen"}
V1 = {}  # round -> world -> the same measures (no live world: one post per screen, no follow-ups)
v1rows = {}
for rd in ("r101", "r102"):
    rows_, _ = load(os.path.join(HERE, "../../../data/llm_bias/v1_v2_spark"), rd)
    v1rows[rd] = [r for r in rows_ if r["turn"] == "read" and r["outcome"] == "chose" and not r["draw"]]
    by = defaultdict(list)
    for r in v1rows[rd]:
        by[(S[r["model"]], S[r["posts_by"]])].append(r)
    V1[rd] = {k: {"posts read": len(rs), "actions chosen": sum(map(A.n_engage, rs)), "upvoted": sum(A.has(r, "like_post") for r in rs),
                  "downvoted": sum(A.has(r, "dislike_post") for r in rs), "commented": sum(A.has(r, "create_comment") for r in rs),
                  "did nothing": sum(not A.did_anything(r) for r in rs),
                  **Counter(a["action"] for r in rs for a in r["actions"])} for k, rs in by.items()}
ACTS = sorted(set(ACTS) | {k for w in V1.values() for c in w.values() for k in c} - set(COUNTS))
MEAS = [("posts read", "posts read"), ("actions chosen", "total engagement (every action but do nothing)"),
        ("upvoted", "posts upvoted"), ("downvoted", "posts downvoted"), ("commented", "posts commented on"),
        ("did nothing", "posts where it only did nothing")] + [(a, "action: " + a) for a in ACTS] + [
        ("follow-up screens", "follow-up screens opened (live world only; not part of engagement)")]


def long_sheet(title, note, sources):
    """One row per world per measure; for each source (name, world -> counts) the count and the rate per 100 posts read."""
    head = ["reader", "posts by", "measure"] + [f"{n}: {c}" for n, _ in sources for c in ("count", "per 100 posts read")]
    if len(sources) >= 2:
        head += [f"change per 100: {sources[-1][0]} - {sources[-2][0]}"]
    body = []
    for m, pb in WORLDS:
        for key, label in MEAS:
            row, rates = [m, pb, label], []
            for _, got in sources:
                g = got.get((m, pb))
                v = (g.get(key) or 0) if g else None
                rate = None if g is None or key == "posts read" else round(100 * v / g["posts read"], 1)
                row += [v, rate]; rates.append(rate)
            if len(sources) >= 2:
                row.append(round(rates[-1] - rates[-2], 1) if None not in rates[-2:] else None)
            if any(x for x in row[3:]):
                body.append(row)
    sheet(title, head, body, {1: 9, 2: 9, 3: 40}, note=note, base_col=(0, 1))


at = lambda tk, t: {(a, b): r for (a, b, tt), r in tk.items() if tt == t}  # noqa: E731
long_sheet("1 No timeline", "NO TIMELINE. V1 (rounds 101, 102: one post per screen, never any comments) next to each V3 part's "
           "tick 0 (each user's first post, before anyone commented; 100 posts read per world). Baselines shaded.",
           [("V1 r101", V1["r101"]), ("V1 r102", V1["r102"])] + [(f"V3 {p} tick 0", at(tk, 0)) for p, tk in zip(PART, tick)])
long_sheet("2 With timeline", "WITH THE TIMELINE: each V3 part over the whole scroll (every tick, comments pile up), parts kept "
           "apart, then the change. Engagement = actions on the post screen (as in V1); follow-up screens are a row of their own.",
           [(f"V3 {p} all ticks", at(tk, "all")) for p, tk in zip(PART, tick)])

# own-AI bias, every set (analyze_v2's double difference and 95% range, as in every analysis)
sets = {f"V1 {rd}": [{"model": S[r["model"]], "posts_by": S[r["posts_by"]], "post_key": r["post_key"], "user_id": r["user_id"],
                      "eng": A.n_engage(r), "up": int(A.has(r, "like_post"))} for r in rs] for rd, rs in v1rows.items()}
for p, x in zip(PART, data):
    rs = [{"model": r["reader"], "posts_by": r["posts_by"], "post_key": r["post_key"], "user_id": r["user_id"], "tick": r["tick"],
           "eng": r["n_chosen"], "up": r["upvoted"]} for r in x["reads"]]
    sets[f"V3 {p} tick 0 (no timeline)"] = [r for r in rs if r["tick"] == 0]
    sets[f"V3 {p} all ticks (timeline)"] = rs
body = []
for name, rs in sets.items():
    for i, j in (("gemma", "llama"), ("gemma", "qwen"), ("llama", "qwen")):
        row = [name, f"{i} vs {j}"]
        for k in ("eng", "up"):
            f = lambda r, k=k: 100 * r[k]  # noqa: E731
            v, (lo, hi) = A.dd(rs, i, j, f), A.boot(rs, i, j, f, n=300)
            row += [round(v, 1), round(lo, 1), round(hi, 1), "CLEAR" if lo > 0 or hi < 0 else "no"]
        body.append(row)
sheet("3 Own-AI bias", ["set", "pair", "engagement bias", "low", "high", "clear?", "upvote bias", "low", "high", "clear?"], body,
      {1: 34, 2: 16}, note="Own-AI bias = (i on i's posts - j on i's posts) - (i on j's posts - j on j's posts), per 100 posts read. "
      "+ = both AIs give their own AI's posts more. 95% range from resampling posts and users (300 draws).")

# one sheet per part per world, tick by tick, every action
for p, tk in zip(PART, tick):
    for m, pb in WORLDS:
        ticks = sorted(t for (a, b, t) in tk if (a, b) == (m, pb) and t != "all") + ["all"]
        acts = [a for a in ACTS if any(tk[(m, pb, t)].get(a) for t in ticks)]
        body = [[t] + [tk[(m, pb, t)][c] for c in COUNTS + RATES] + [tk[(m, pb, t)].get(a) or 0 for a in acts] for t in ticks]
        sheet(f"P{p[-1]} {m} reads {pb}" + (" (base)" if m == pb else ""), ["tick"] + COUNTS + RATES + acts, body, {1: 8},
              note=f"{p.upper()}: {m} plays all 100 users, reading the {sum(1 for x in data[0]['posts'] if x['ai'] == pb)} posts {pb} wrote"
                   + (" (baseline)" if m == pb else "") + ". One row per tick (tick 0 = no timeline yet); last row = all ticks. "
                   "Action columns = times chosen on the post screen at that tick.")

# Worlds: each world all together (whole scroll)
head = ["reader", "posts by"] + [f"{p}: {c}" for p in PART for c in COUNTS + RATES]
if len(DIRS) > 1:
    head += [f"change: {c}" for c in RATES]
body = []
for m, pb in WORLDS:
    got = [tk[(m, pb, "all")] for tk in tick]
    row = [m, pb] + [g[c] for g in got for c in COUNTS + RATES]
    if len(DIRS) > 1:
        row += [round(got[1][c] - got[0][c], 1) for c in RATES]
    body.append(row)
sheet("Worlds", head, body, {1: 9, 2: 9}, note="Each world all together (baselines first, shaded).", base_col=(0, 1))

# Actions
acts = sorted({r["action"] for x in data for r in x["by_cell_action"]})
act = [{(r["reader"], r["posts_by"], r["action"]): r for r in x["by_cell_action"]} for x in data]
head = ["reader", "posts by", "action"] + [f"{p}: {c}" for p in PART for c in ("chosen", "done", "failed", "chosen per 100", "done per 100")]
body = []
for m, pb in WORLDS:
    for a in acts:
        row = [m, pb, a]
        for ak in act:
            r = ak.get((m, pb, a), {"chosen": 0, "done": 0, "chosen_per_100": 0, "done_per_100": 0})
            row += [r["chosen"], r["done"], r["chosen"] - r["done"], r["chosen_per_100"], r["done_per_100"]]
        body.append(row)
sheet("Actions", head, body, {1: 9, 2: 9, 3: 22}, note="Every action in every world. Filter 'action' to compare one action across worlds.",
      base_col=(0, 1))

# How hard each action is, and how often each AI picks it (per reader AI, its 3 worlds together)
for p, x in zip(PART, data):
    t = x["effort_tiers"]
    sheet("Effort tiers" + (f" {p}" if len(DIRS) > 1 else ""), list(t[0]), [list(r.values()) for r in t], {3: 50},
          note="Each AI's actions (do nothing left out) by how much the action asks: 1 one click ... 5 go to another "
               "screen first, then write. Do the AIs pick the easy ones?")
    e = x["action_effort"]
    sheet("Action effort" + (f" {p}" if len(DIRS) > 1 else ""), list(e[0]), [list(r.values()) for r in e], {2: 20, 4: 40},
          note="Every action each AI chose, on every screen. Answer tokens / seconds: answers that held only this "
               "action (seconds = time for one answer while 8 run at once; compare within one AI).")

# Posts: each post, and what each world did with it
for p, x in zip(PART, data):
    by = defaultdict(Counter)
    for r in x["reads"]:
        c = by[(r["post_key"], r["reader"])]
        c["read"] += 1; c["up"] += r["upvoted"]; c["down"] += r["downvoted"]; c["com"] += r["commented"]; c["act"] += r["n_chosen"]
    readers = sorted({r["reader"] for r in x["reads"]})
    head = ["post", "written by AI", "writer user", "topic", "characters"] + [
        f"{m}: {c}" for m in readers for c in ("times read", "upvoted", "downvoted", "commented", "actions", "upvoted %")]
    body = []
    for q in x["posts"]:
        row = [q["post_key"], q["ai"], q["user_id"], q["topic"], q["chars"]]
        for m in readers:
            c = by[(q["post_key"], m)]
            row += [c["read"], c["up"], c["down"], c["com"], c["act"], round(100 * c["up"] / c["read"], 1) if c["read"] else None]
        body.append(row)
    sheet("Posts" + (f" {p}" if len(DIRS) > 1 else ""), head, body, {1: 24, 4: 16},
          note="Each post (the same posts in both parts). Every post is read by 99 users in each of the 3 reader worlds.")

# Users: each reader user in each world
for p, x in zip(PART, data):
    by = defaultdict(Counter)
    for r in x["reads"]:
        c = by[(r["reader"], r["posts_by"], r["user_id"])]
        c["read"] += 1; c["act"] += r["n_chosen"]; c["done"] += r["n_done"]; c["up"] += r["upvoted"]; c["down"] += r["downvoted"]
        c["com"] += r["commented"]; c["fu"] += r["followup_screens"]
    head = ["reader", "posts by", "user", "posts read", "actions chosen", "actions done", "upvoted", "downvoted", "commented",
            "follow-up screens", "engagement per 100", "upvoted %"]
    body = [[m, pb, u, c["read"], c["act"], c["done"], c["up"], c["down"], c["com"], c["fu"],
             round(100 * c["act"] / c["read"], 1), round(100 * c["up"] / c["read"], 1)]
            for (m, pb, u), c in sorted(by.items(), key=lambda kv: (WORLDS.index(kv[0][:2]), kv[0][2]))]
    sheet("Users" + (f" {p}" if len(DIRS) > 1 else ""), head, body, {1: 9, 2: 9}, base_col=(0, 1),
          note="Each user (the same 100 people in every world, each played by the world's reader AI).")

# every row
for p, x in zip(PART, data):
    for name, key in (("Reads", "reads"), ("Follow-ups", "followups")):
        rs = x[key]
        sheet(name + (f" {p}" if len(DIRS) > 1 else ""), list(rs[0]), [list(r.values()) for r in rs], {5: 24, 12: 30, 13: 20})

for i, name in enumerate(["Read me", "1 No timeline", "2 With timeline", "3 Own-AI bias"]):
    wb.move_sheet(name, i - wb.sheetnames.index(name))
wb.save(OUT)
print(f"wrote {OUT}: {len(wb.sheetnames)} sheets: {', '.join(wb.sheetnames)}")
