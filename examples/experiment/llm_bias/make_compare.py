"""Compare every round of the Spark study (101, 102, 103, 104) and the two versions (V1 = 101+102, V2 = 103+104), for
every measure, in one JSON file (for the Inside page) and one Excel workbook (log 14.37).

    oasis-env/bin/python examples/experiment/llm_bias/make_compare.py data/llm_bias/v2_spark ~/Desktop/LLM_Bias

IN PLAIN WORDS
--------------
Views: R101, R102, R103, R104, V1 (101+102 pooled), V2 (103+104 pooled). Pooling = every screen counts once.
Measures, main one first: total engagement (every action except "do nothing", per 100 screens), engaged (% of screens
with any action), did nothing (%), then each of the 27 actions (% of screens where the user took it).
For every view x measure: the 3x3 grid (reader AI x whose posts) and the own-AI bias for each pair of AIs
  (i reading i - j reading i) - (i reading j - j reading j)          (same double difference as analyze_v2.dd)
Changes: V2 - V1, and the round-to-round wobble inside each version (R102 - R101, R104 - R103).
95% ranges: 1,000 resamples of users and posts together. The 100 users are the same people in every round, so one
user draw is shared by all rounds; post keys are unique per round. A cell's mean in a draw is
  sum(w_post * w_user * value) / sum(w_post * w_user * screens)
which is exactly what analyze_v2.boot does by copying rows, just done with numpy.
Also: posting turn, noise floor, by stance, by topic and unreadable screens, per round.
Writes <out>/compare.json and <out>/LLM_Bias_Compare.xlsx.
"""
import json, os, sys
from collections import Counter, defaultdict

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from make_report_v2 import MODELS, SHORT, STANCES, analyze_topic, lbl, load, post_index  # noqa: E402
from run_v2 import ACTION_NAMES, TOPICS  # noqa: E402

ROUNDS = [101, 102, 103, 104]
VIEWS = {"R101": [101], "R102": [102], "R103": [103], "R104": [104], "V1": [101, 102], "V2": [103, 104]}
VIEW_LABEL = {"R101": "Round 101 (V1)", "R102": "Round 102 (V1)", "R103": "Round 103 (V2)", "R104": "Round 104 (V2)",
              "V1": "V1: rounds 101 + 102", "V2": "V2: rounds 103 + 104"}
CONTRASTS = {"V2-V1": ("V2", "V1"), "R102-R101": ("R102", "R101"), "R104-R103": ("R104", "R103")}
CONTRAST_LABEL = {"V2-V1": "V2 − V1 (does knowing its name change it?)", "R102-R101": "R102 − R101 (wobble inside V1)",
                  "R104-R103": "R104 − R103 (wobble inside V2)"}
METRICS = ([("tot", "Total engagement", "actions per 100 screens"), ("any", "Engaged (any action)", "% of screens"),
            ("none", "Did nothing", "% of screens")] +
           [("a:" + a, lbl(a), "% of screens") for a in ACTION_NAMES])
PAIRS = [(MODELS[a], MODELS[b]) for a in range(3) for b in range(a + 1, 3)]
B, SEED = 1000, 7


def values(r):
    """One screen -> every measure, x100 (so means are per 100 screens / percent)."""
    acts = {x["action"] for x in r["actions"]}
    tot = sum(x["action"] != "do_nothing" for x in r["actions"])
    return [100 * tot, 100 * (tot > 0), 100 * (tot == 0)] + [100 * (a in acts) for a in ACTION_NAMES]


def lohi(d):
    d = d[np.isfinite(d)]
    return (float(np.percentile(d, 2.5)), float(np.percentile(d, 97.5))) if len(d) > 20 else (None, None)


def bias_of(g, i, j):
    """g[(reader, poster)] -> array (any shape); the double difference."""
    return (g[(i, i)] - g[(j, i)]) - (g[(i, j)] - g[(j, j)])


def main(res, out):
    os.makedirs(out, exist_ok=True)
    reads, post_rows = load(res)
    reads = [r for r in reads if r["round"] in ROUNDS]
    post_rows = [r for r in post_rows if r["round"] in ROUNDS]
    R = [r for r in reads if r["outcome"] == "chose" and not r["draw"]]
    users = sorted({r["user_id"] for r in R})
    posts = sorted({r["post_key"] for r in R})
    ui, pi = {u: k for k, u in enumerate(users)}, {p: k for k, p in enumerate(posts)}
    K = len(METRICS)

    # ---- per (round, reader, poster): sums per (post, user) pair, then the bootstrap numerators/denominators
    by = defaultdict(list)
    for r in R:
        by[(r["round"], r["model"], r["posts_by"])].append(r)
    rng = np.random.default_rng(SEED)
    wu = rng.multinomial(len(users), np.full(len(users), 1 / len(users)), size=B).astype(np.float64)
    wp = rng.multinomial(len(posts), np.full(len(posts), 1 / len(posts)), size=B).astype(np.float64)
    num0, den0, numb, denb = {}, {}, {}, {}
    for key, rows in by.items():
        pair = np.array([pi[r["post_key"]] * len(users) + ui[r["user_id"]] for r in rows])
        V = np.array([values(r) for r in rows], dtype=np.float64)
        uniq, inv = np.unique(pair, return_inverse=True)
        S = np.zeros((len(uniq), K))
        np.add.at(S, inv, V)
        cnt = np.bincount(inv).astype(np.float64)
        W = wp[:, uniq // len(users)] * wu[:, uniq % len(users)]
        num0[key], den0[key], numb[key], denb[key] = V.sum(0), float(len(rows)), W @ S, W @ cnt

    def cells(view, boot):
        """(reader, poster) -> mean per measure: shape (K,) for the point estimate, (B, K) for the draws."""
        g = {}
        for m in MODELS:
            for pb in MODELS:
                ks = [(rd, m, pb) for rd in VIEWS[view] if (rd, m, pb) in num0]
                if boot:
                    g[(m, pb)] = sum(numb[k] for k in ks) / sum(denb[k] for k in ks)[:, None]
                else:
                    g[(m, pb)] = sum(num0[k] for k in ks) / sum(den0[k] for k in ks)
        return g

    P = {v: cells(v, False) for v in VIEWS}
    D = {v: cells(v, True) for v in VIEWS}
    N = {v: {(m, pb): int(sum(den0.get((rd, m, pb), 0) for rd in VIEWS[v])) for m in MODELS for pb in MODELS} for v in VIEWS}

    grid, bias = {}, {}
    for v in VIEWS:
        grid[v], bias[v] = {}, {}
        for k, (mk, _, _) in enumerate(METRICS):
            grid[v][mk] = {f"{m}|{pb}": [float(P[v][(m, pb)][k]), *lohi(D[v][(m, pb)][:, k]), N[v][(m, pb)]]
                           for m in MODELS for pb in MODELS}
            bias[v][mk] = [{"i": i, "j": j, "v": float(bias_of(P[v], i, j)[k]), **dict(zip(("lo", "hi"), lohi(bias_of(D[v], i, j)[:, k])))}
                           for i, j in PAIRS]
    contrast = {}
    for c, (a, b) in CONTRASTS.items():
        contrast[c] = {"grid": {}, "bias": {}}
        for k, (mk, _, _) in enumerate(METRICS):
            contrast[c]["grid"][mk] = {f"{m}|{pb}": [float(P[a][(m, pb)][k] - P[b][(m, pb)][k]),
                                                    *lohi(D[a][(m, pb)][:, k] - D[b][(m, pb)][:, k])] for m in MODELS for pb in MODELS}
            contrast[c]["bias"][mk] = [{"i": i, "j": j, "v": float(bias_of(P[a], i, j)[k] - bias_of(P[b], i, j)[k]),
                                        **dict(zip(("lo", "hi"), lohi(bias_of(D[a], i, j)[:, k] - bias_of(D[b], i, j)[:, k])))}
                                       for i, j in PAIRS]

    # ---- posting turn, noise floor, unreadable: per round
    posting = {}
    for rd in ROUNDS:
        posting[rd] = {}
        for m in MODELS:
            ps = [r for r in post_rows if r["round"] == rd and r["model"] == m]
            wrote = [x for r in ps for x in r["actions"] if x["action"] == "create_post"]
            posting[rd][m] = {"users": len(ps), "posters": sum(any(x["action"] == "create_post" for x in r["actions"]) for r in ps),
                              "posts": len(wrote), "unreadable": sum(r["outcome"] != "chose" for r in ps),
                              "topics": dict(Counter(analyze_topic(x.get("subreddit")) or "off-topic" for x in wrote)),
                              "other": dict(Counter(x["action"] for r in ps for x in r["actions"] if x["action"] != "create_post").most_common())}
    first = {(r["round"], r["model"], r["user_id"], r["post_key"]): r for r in R}
    noise = {}
    for rd in ROUNDS:
        noise[rd] = {}
        for m in MODELS:
            prs = [(first.get((rd, m, r["user_id"], r["post_key"])), r) for r in reads
                   if r["round"] == rd and r["draw"] and r["model"] == m and r["outcome"] == "chose"]
            prs = [(values(x), values(y)) for x, y in prs if x]
            if prs:
                n = len(prs)
                noise[rd][m] = {"n": n, "same_eng": 100 * sum(x[1] == y[1] for x, y in prs) / n,
                                "same_n": 100 * sum(x[0] == y[0] for x, y in prs) / n,
                                "same_up": 100 * sum(x[3] == y[3] for x, y in prs) / n,
                                "tot_a": sum(x[0] for x, _ in prs) / n, "tot_b": sum(y[0] for _, y in prs) / n,
                                "up_a": sum(x[3] for x, _ in prs) / n, "up_b": sum(y[3] for _, y in prs) / n}
    unread = {}
    for rd in ROUNDS:
        unread[rd] = {}
        for m in MODELS:
            for pb in MODELS:
                rs = [r for r in reads if r["round"] == rd and r["model"] == m and r["posts_by"] == pb and not r["draw"]]
                bad = sum(r["outcome"] != "chose" for r in rs)
                unread[rd][f"{m}|{pb}"] = {"screens": len(rs), "unreadable": bad, "pct": 100 * bad / len(rs) if rs else None}

    # ---- by stance and by topic (points only): tot, any, upvote
    split = {}
    for name, keyf, order in (("stance", lambda r: r["stance"], STANCES), ("topic", lambda r: r["topic"], list(TOPICS))):
        acc = defaultdict(lambda: [0, 0, 0, 0])  # screens, tot, any, up (x100 sums)
        for r in R:
            vs = values(r)
            for v in VIEWS:
                if r["round"] in VIEWS[v]:
                    a = acc[(v, r["model"], r["posts_by"], keyf(r))]
                    a[0] += 1; a[1] += vs[0]; a[2] += vs[1]; a[3] += vs[3]
        split[name] = {"order": order, "rows": [{"view": v, "reader": m, "poster": pb, "group": g, "n": a[0],
                                                  "tot": a[1] / a[0], "any": a[2] / a[0], "up": a[3] / a[0]}
                                                 for (v, m, pb, g), a in sorted(acc.items(), key=lambda kv: (list(VIEWS).index(kv[0][0]),
                                                 MODELS.index(kv[0][1]), MODELS.index(kv[0][2]), order.index(kv[0][3]) if kv[0][3] in order else 99))]}

    J = {"models": MODELS, "short": SHORT, "metrics": [{"key": k, "name": n, "unit": u} for k, n, u in METRICS],
         "views": {v: {"label": VIEW_LABEL[v], "rounds": VIEWS[v], "screens": sum(N[v].values())} for v in VIEWS},
         "contrasts": {c: {"label": CONTRAST_LABEL[c], "a": a, "b": b} for c, (a, b) in CONTRASTS.items()},
         "grid": grid, "bias": bias, "contrast": contrast, "posting": posting, "noise": noise, "unreadable": unread,
         "split": split, "boot": {"draws": B, "seed": SEED, "users": len(users), "posts": len(posts)}}
    jp = os.path.join(out, "compare.json")
    json.dump(J, open(jp, "w"), indent=0)
    xp = workbook(J, os.path.join(out, "LLM_Bias_Compare.xlsx"))
    print(f"wrote {jp} ({os.path.getsize(jp) / 1e3:.0f} kB) and {xp}")
    for v in ("V1", "V2"):
        print(v, {f"{SHORT[b['i']]}-{SHORT[b['j']]}": [round(b[x], 1) for x in ("v", "lo", "hi")] for b in bias[v]["tot"]})
    print("V2-V1", {f"{SHORT[b['i']]}-{SHORT[b['j']]}": [round(b[x], 1) for x in ("v", "lo", "hi")] for b in contrast["V2-V1"]["bias"]["tot"]})


# ================================================================ workbook
HEAD = PatternFill("solid", fgColor="13212E")
DIAG = PatternFill("solid", fgColor="F3F6F8")
OWN = PatternFill("solid", fgColor="CFE4FA")    # range above 0: own AI's posts favoured / V2 higher
OTHER = PatternFill("solid", fgColor="FBD9C6")  # range below 0: other AI's posts favoured / V2 lower
WHITE, BOLD, TITLE = Font(bold=True, color="FFFFFF"), Font(bold=True), Font(bold=True, size=14)
NUM = "0.0"
SGN = "+0.0;-0.0;0.0"


def verdict(lo, hi, kind="bias"):
    if lo is None:
        return "-"
    if kind == "bias":
        return "own AI's posts favoured" if lo > 0 else "other AI's posts favoured" if hi < 0 else "no clear bias"
    return "clearly higher" if lo > 0 else "clearly lower" if hi < 0 else "no clear change"


def fill_for(lo, hi):
    return OWN if lo is not None and lo > 0 else OTHER if hi is not None and hi < 0 else None


class Sheet:
    """A tiny row writer: keeps the current row, styles headers, remembers column widths."""
    def __init__(self, wb, name, widths):
        self.ws, self.r = wb.create_sheet(name), 1
        for k, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(k)].width = w

    def row(self, vals, font=None, fill=None, fmt=None, fills=None, wrap=False):
        for k, v in enumerate(vals, 1):
            c = self.ws.cell(self.r, k, v)
            if font: c.font = font
            if fill: c.fill = fill
            if fills and fills.get(k): c.fill = fills[k]
            if isinstance(v, float) and fmt: c.number_format = fmt.get(k, NUM) if isinstance(fmt, dict) else fmt
            if wrap: c.alignment = Alignment(wrap_text=True, vertical="top")
        self.r += 1

    def head(self, vals):
        self.row(vals, font=WHITE, fill=HEAD)
        for k in range(1, len(vals) + 1):
            self.ws.cell(self.r - 1, k).alignment = Alignment(wrap_text=True, vertical="center")

    def title(self, text, size=14):
        self.ws.cell(self.r, 1, text).font = Font(bold=True, size=size)
        self.r += 1

    def note(self, text):
        self.ws.cell(self.r, 1, text).font = Font(italic=True, color="566473")
        self.r += 1

    def gap(self, n=1):
        self.r += n


def workbook(J, path):
    S, M = J["short"], J["models"]
    metrics = J["metrics"]
    wb = Workbook()
    wb.remove(wb.active)

    # README
    s = Sheet(wb, "README", [110])
    s.title("LLM Bias: every round and V1 vs V2, compared")
    for line in [
        "V1 = rounds 101 + 102: each AI is given only the person to play.",
        "V2 = rounds 103 + 104: exactly the same, plus one line before the person: \"I am Qwen3 8B, an AI model made by Alibaba, "
        "simulating this profile:\" (Llama 3.1 8B / Meta, Gemma 3 12B / Google). The AI is still never told who wrote a post.",
        "Question: does knowing which model it is change how much an AI favours posts written by its own AI?",
        "",
        "HOW TO READ",
        "Grid = reader AI (rows) x whose posts it read (columns). The diagonal (same AI wrote and read) is the baseline.",
        "Total engagement (the main measure) = every action on a screen except 'do nothing', per 100 screens. Upvote + comment = 2.",
        "Engaged = % of screens with at least one action. Did nothing = % of screens with no action. Each action = % of screens where the user took it.",
        "Own-AI bias for AIs i and j = (i reading i's posts - j reading i's posts) - (i reading j's posts - j reading j's posts).",
        "  It is the same number whichever AI you call i. Above 0 = in this pair, posts get more of the measure from the same AI that",
        "  wrote them (beyond one AI simply being more active); below 0 = posts get more from the other AI.",
        "Change V2 - V1 = the V2 number minus the V1 number. R102 - R101 and R104 - R103 show how much a number wobbles between two",
        "  rounds of the SAME version, so a real V1-vs-V2 change should be bigger than those.",
        f"95% range: {J['boot']['draws']:,} resamples of the {J['boot']['users']} users and {J['boot']['posts']} posts together (seed {J['boot']['seed']}). "
        "The same 100 users are in every round, so each resample uses the same user draw for all rounds.",
        "Blue cell = the 95% range is entirely above 0 (own AI's posts favoured / clearly higher). Orange = entirely below 0. Uncoloured = range crosses 0 (no clear result).",
        "Pooled views (V1, V2) count every screen once. Screens with an unreadable answer and the noise-floor re-reads are left out.",
        "",
        "SHEETS",
        "Total engagement: the main measure, round by round and V1 | V2 | V2 - V1, with the bias tables.",
        "V1 vs V2: every measure, one block each (grid + bias, with the change and its range).",
        "Round by round: every measure, one block each, rounds 101-104 side by side.",
        "All grids / All biases: every number above in one long table each, with filters (pick a view and a measure).",
        "Posting turn, Noise floor, By stance, By topic, Unreadable: per round.",
        "",
        "Made by examples/experiment/llm_bias/make_compare.py from data/llm_bias/v2_spark (log 14.37-14.38).",
    ]:
        s.row([line], font=BOLD if line.isupper() and line else None)

    def view_grid(s, v, mk, col0=1):
        """Write one 4x4 grid starting at column col0 on the current row; returns rows used."""
        g = J["grid"][v][mk]
        r0 = s.r
        ws = s.ws
        ws.cell(r0, col0, VIEW_LABEL[v]).font = BOLD
        hdr = ["Reader ↓ / posts by →"] + [S[m] for m in M]
        for k, h in enumerate(hdr):
            c = ws.cell(r0 + 1, col0 + k, h)
            c.font, c.fill = WHITE, HEAD
        for a, m in enumerate(M):
            ws.cell(r0 + 2 + a, col0, S[m]).font = BOLD
            for b, pb in enumerate(M):
                c = ws.cell(r0 + 2 + a, col0 + 1 + b, round(g[f"{m}|{pb}"][0], 4))
                c.number_format = NUM
                if m == pb: c.fill = DIAG
        return 5

    def contrast_grid(s, c, mk, col0):
        g = J["contrast"][c]["grid"][mk]
        r0, ws = s.r, s.ws
        ws.cell(r0, col0, c.replace("-", " − ")).font = BOLD
        for k, h in enumerate(["Reader ↓ / posts by →"] + [S[m] for m in M]):
            cc = ws.cell(r0 + 1, col0 + k, h)
            cc.font, cc.fill = WHITE, HEAD
        for a, m in enumerate(M):
            ws.cell(r0 + 2 + a, col0, S[m]).font = BOLD
            for b, pb in enumerate(M):
                v, lo, hi = g[f"{m}|{pb}"]
                cc = ws.cell(r0 + 2 + a, col0 + 1 + b, round(v, 4))
                cc.number_format = SGN
                f = fill_for(lo, hi)
                if f: cc.fill = f
        return 5

    def bias_table(s, mk, views, contrasts=()):
        s.head(["Pair (i vs j)"] + [x for v in views for x in (f"{v} bias", "low", "high")] +
               [x for c in contrasts for x in (c.replace("-", " − "), "low", "high")] + ["Reading"])
        for p in range(3):
            vals, fills, fmt = [f"{S[PAIRS[p][0]]} vs {S[PAIRS[p][1]]}"], {}, {}
            for v in views:
                b = J["bias"][v][mk][p]
                fills[len(vals) + 1] = fill_for(b["lo"], b["hi"])
                vals += [b["v"], b["lo"], b["hi"]]
            for c in contrasts:
                b = J["contrast"][c]["bias"][mk][p]
                fills[len(vals) + 1] = fill_for(b["lo"], b["hi"])
                vals += [b["v"], b["lo"], b["hi"]]
            last = J["contrast"][contrasts[0]]["bias"][mk][p] if contrasts else J["bias"][views[-1]][mk][p]
            vals.append(", ".join(f"{v}: {verdict(J['bias'][v][mk][p]['lo'], J['bias'][v][mk][p]['hi'])}" for v in views) +
                        (f"; change: {verdict(last['lo'], last['hi'], 'change')}" if contrasts else ""))
            s.row(vals, fmt=SGN, fills=fills)

    # Total engagement
    s = Sheet(wb, "Total engagement", [22 if k % 5 == 0 else 12 for k in range(20)])
    s.title("Total engagement: every action except 'do nothing', per 100 screens")
    s.note("Rows = the AI playing the reader, columns = the AI that wrote the posts. Diagonal (grey) = baseline.")
    s.gap()
    s.title("Round by round", 12)
    for k, v in enumerate(["R101", "R102", "R103", "R104"]):
        view_grid(s, v, "tot", 1 + 5 * k)
    s.gap(6)
    s.title("V1 | V2 | V2 − V1 (blue = clearly higher in V2, orange = clearly lower)", 12)
    view_grid(s, "V1", "tot", 1); view_grid(s, "V2", "tot", 6); contrast_grid(s, "V2-V1", "tot", 11)
    s.gap(6)
    s.title("Own-AI bias on total engagement, per round", 12)
    bias_table(s, "tot", ["R101", "R102", "R103", "R104"])
    s.gap()
    s.title("Own-AI bias on total engagement: V1 vs V2", 12)
    bias_table(s, "tot", ["V1", "V2"], ["V2-V1", "R102-R101", "R104-R103"])
    s.ws.freeze_panes = "A4"

    # V1 vs V2: every measure
    s = Sheet(wb, "V1 vs V2", [24, 18] + [11] * 14 + [90])
    s.title("V1 vs V2, every measure (total engagement first, then each action)")
    s.note("Each block: the 9 grid cells (V1, V2, change with 95% range), then the own-AI bias for each pair. Blue/orange = range excludes 0.")
    s.ws.freeze_panes = "A3"
    for mt in metrics:
        mk = mt["key"]
        s.gap()
        s.row([f"{mt['name']} ({mt['unit']})"], font=TITLE)
        s.head(["Reader AI", "Posts by", "V1", "V1 screens", "V2", "V2 screens", "V2 − V1", "low", "high", "Reading"])
        for m in M:
            for pb in M:
                k = f"{m}|{pb}"
                a, b, c = J["grid"]["V1"][mk][k], J["grid"]["V2"][mk][k], J["contrast"]["V2-V1"]["grid"][mk][k]
                s.row([S[m], S[pb] + (" (baseline)" if m == pb else ""), a[0], a[3], b[0], b[3], c[0], c[1], c[2], verdict(c[1], c[2], "change")],
                      fmt={3: NUM, 5: NUM, 7: SGN, 8: SGN, 9: SGN}, fills={7: fill_for(c[1], c[2])})
        s.gap()
        bias_table(s, mk, ["V1", "V2"], ["V2-V1", "R102-R101", "R104-R103"])

    # Round by round: every measure
    s = Sheet(wb, "Round by round", [24, 18] + [11] * 11 + [110])
    s.title("Rounds 101-104 side by side, every measure")
    s.note("Each block: the 9 grid cells per round, then the own-AI bias per round with its 95% range.")
    s.ws.freeze_panes = "A3"
    for mt in metrics:
        mk = mt["key"]
        s.gap()
        s.row([f"{mt['name']} ({mt['unit']})"], font=TITLE)
        s.head(["Reader AI", "Posts by", "R101", "R102", "R103", "R104", "V1 (101+102)", "V2 (103+104)"])
        for m in M:
            for pb in M:
                k = f"{m}|{pb}"
                s.row([S[m], S[pb] + (" (baseline)" if m == pb else "")] + [J["grid"][v][mk][k][0] for v in ("R101", "R102", "R103", "R104", "V1", "V2")],
                      fmt=NUM)
        s.gap()
        bias_table(s, mk, ["R101", "R102", "R103", "R104"])

    # All grids (long)
    s = Sheet(wb, "All grids", [20, 26, 22, 12, 12, 12, 12, 12, 12, 20])
    s.head(["View", "Measure", "Unit", "Reader AI", "Posts by", "Baseline?", "Screens", "Value", "low", "high"])
    for v in VIEWS:
        for mt in metrics:
            for m in M:
                for pb in M:
                    x = J["grid"][v][mt["key"]][f"{m}|{pb}"]
                    s.row([v, mt["name"], mt["unit"], S[m], S[pb], "yes" if m == pb else "no", x[3], x[0], x[1], x[2]], fmt=NUM)
    for c in CONTRASTS:
        for mt in metrics:
            for m in M:
                for pb in M:
                    x = J["contrast"][c]["grid"][mt["key"]][f"{m}|{pb}"]
                    s.row([c.replace("-", " − "), mt["name"], mt["unit"], S[m], S[pb], "yes" if m == pb else "no", None, x[0], x[1], x[2]],
                          fmt=SGN, fills={8: fill_for(x[1], x[2])})
    s.ws.freeze_panes = "A2"
    s.ws.auto_filter.ref = f"A1:J{s.r - 1}"

    # All biases (long)
    s = Sheet(wb, "All biases", [20, 26, 22, 10, 10, 12, 12, 12, 24])
    s.head(["View", "Measure", "Unit", "i", "j", "Bias", "low", "high", "Reading"])
    for v in VIEWS:
        for mt in metrics:
            for b in J["bias"][v][mt["key"]]:
                s.row([v, mt["name"], mt["unit"], S[b["i"]], S[b["j"]], b["v"], b["lo"], b["hi"], verdict(b["lo"], b["hi"])],
                      fmt=SGN, fills={6: fill_for(b["lo"], b["hi"])})
    for c in CONTRASTS:
        for mt in metrics:
            for b in J["contrast"][c]["bias"][mt["key"]]:
                s.row([c.replace("-", " − "), mt["name"], mt["unit"], S[b["i"]], S[b["j"]], b["v"], b["lo"], b["hi"], verdict(b["lo"], b["hi"], "change")],
                      fmt=SGN, fills={6: fill_for(b["lo"], b["hi"])})
    s.ws.freeze_panes = "A2"
    s.ws.auto_filter.ref = f"A1:I{s.r - 1}"

    # Posting turn
    s = Sheet(wb, "Posting turn", [10, 12, 9, 15, 9, 12] + [12] * len(TOPICS) + [12, 70])
    s.title("Posting turn: empty feed, every action available, nobody told to post")
    s.head(["Round", "AI", "Users", "Users who posted", "Posts", "Unreadable"] + [TOPICS[t][0] for t in TOPICS] + ["off-topic", "Other actions taken"])
    for rd in ROUNDS:
        for m in M:
            p = J["posting"][rd][m]
            s.row([rd, S[m], p["users"], p["posters"], p["posts"], p["unreadable"]] + [p["topics"].get(t, 0) for t in TOPICS] +
                  [p["topics"].get("off-topic", 0), ", ".join(f"{lbl(a)} {n}" for a, n in p["other"].items())])
    s.ws.freeze_panes = "A3"

    # Noise floor
    s = Sheet(wb, "Noise floor", [10, 12, 12, 16, 16, 16, 14, 14, 14, 14])
    s.title("Noise floor: the same AI re-reads the same post as the same user, fresh randomness")
    s.note("A real bias has to be bigger than how much the same AI disagrees with itself.")
    s.head(["Round", "AI", "Screens", "Same engaged / not (%)", "Same number of actions (%)", "Same upvote decision (%)",
            "Total eng. first", "Total eng. re-read", "Upvote % first", "Upvote % re-read"])
    for rd in ROUNDS:
        for m in M:
            n = J["noise"][rd].get(m)
            if n:
                s.row([rd, S[m], n["n"], n["same_eng"], n["same_n"], n["same_up"], n["tot_a"], n["tot_b"], n["up_a"], n["up_b"]], fmt=NUM)
    s.ws.freeze_panes = "A4"

    for name, title in (("stance", "By stance: how the reader feels about the post's topic (love .. hate)"),
                        ("topic", "By topic of the post")):
        sp = J["split"][name]
        s = Sheet(wb, "By " + name, [10, 12, 12, 26, 11, 14, 14, 14])
        s.title(title)
        s.head(["View", "Reader AI", "Posts by", name.capitalize(), "Screens", "Total engagement", "Engaged %", "Upvote %"])
        for x in sp["rows"]:
            g = TOPICS[x["group"]][1] if name == "topic" and x["group"] in TOPICS else x["group"]
            s.row([x["view"], S[x["reader"]], S[x["poster"]], g, x["n"], x["tot"], x["any"], x["up"]], fmt=NUM)
        s.ws.freeze_panes = "A3"
        s.ws.auto_filter.ref = f"A2:H{s.r - 1}"

    s = Sheet(wb, "Unreadable", [10, 12, 12, 12, 13, 13])
    s.title("Unreadable answers (left out of every number): the AI's answer could not be read after 3 tries, or timed out")
    s.head(["Round", "Reader AI", "Posts by", "Screens", "Unreadable", "Unreadable %"])
    for rd in ROUNDS:
        for m in M:
            for pb in M:
                u = J["unreadable"][rd][f"{m}|{pb}"]
                s.row([rd, S[m], S[pb], u["screens"], u["unreadable"], u["pct"]], fmt={6: "0.00"})
    s.ws.freeze_panes = "A3"

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 110
    wb.save(path)
    return path


if __name__ == "__main__":
    main(sys.argv[1], os.path.expanduser(sys.argv[2]))
